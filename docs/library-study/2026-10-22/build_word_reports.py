"""Build the two Word exports from their adjacent Markdown sources.

Requires python-docx. Rendering and visual inspection remain separate QA steps.
"""
from pathlib import Path
import re
from docx import Document
from docx.shared import Inches, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn


ROOT=Path(__file__).resolve().parent


def inline(paragraph,text):
    for part in re.split(r'(`[^`]+`|\*\*[^*]+\*\*)',text):
        r=paragraph.add_run(part[1:-1] if part.startswith('`') and part.endswith('`') else part[2:-2] if part.startswith('**') and part.endswith('**') else part)
        if part.startswith('`'):r.font.name='DejaVu Sans Mono';r.font.size=Pt(8.8)
        if part.startswith('**'):r.bold=True


def build(source):
    doc=Document();sec=doc.sections[0]
    sec.page_width=Inches(8.5);sec.page_height=Inches(11)
    sec.top_margin=sec.bottom_margin=Inches(.7);sec.left_margin=sec.right_margin=Inches(.7)
    for name in ['Normal','Title','Subtitle','Heading 1','Heading 2','Heading 3']:
        st=doc.styles[name];st.font.name='Calibri';st.font.color.rgb=RGBColor(0,0,0)
    doc.styles['Normal'].font.size=Pt(10.5)
    doc.styles['Normal'].paragraph_format.space_after=Pt(6)
    doc.styles['Normal'].paragraph_format.line_spacing=1.1
    doc.styles['Title'].font.size=Pt(25)
    doc.styles['Heading 1'].font.size=Pt(16)
    doc.styles['Heading 2'].font.size=Pt(12)
    for name in ['Heading 1','Heading 2','Heading 3']:
        doc.styles[name].paragraph_format.keep_with_next=True
    lines=source.read_text().splitlines();i=0
    while i<len(lines):
        line=lines[i]
        if not line.strip():i+=1;continue
        if line.startswith('```'):
            i+=1
            while i<len(lines) and not lines[i].startswith('```'):
                p=doc.add_paragraph();p.paragraph_format.space_after=Pt(1)
                p.paragraph_format.line_spacing=1
                p.paragraph_format.keep_together=True
                p.paragraph_format.keep_with_next=(i+1<len(lines) and not lines[i+1].startswith('```'))
                r=p.add_run(lines[i]);r.font.name='DejaVu Sans Mono';r.font.size=Pt(9)
                i+=1
            i+=1;continue
        if line.startswith('|'):
            rows=[]
            while i<len(lines) and lines[i].startswith('|'):
                vals=[x.strip() for x in lines[i].strip().strip('|').split('|')]
                if not all(re.fullmatch(r':?-+:?',v or 'x') for v in vals):rows.append(vals)
                i+=1
            table=doc.add_table(rows=0,cols=len(rows[0]));table.autofit=False
            n=len(rows[0]);widths=([2.0,5.1] if n==2 else [1.45,1.65,1.05,2.95] if n==4 else [1.5]+[(7.1-1.5)/(n-1)]*(n-1))
            for col,width in zip(table.columns,widths):col.width=Inches(width)
            for ridx,values in enumerate(rows):
                row=table.add_row()
                pr=row._tr.get_or_add_trPr();avoid=OxmlElement('w:cantSplit');pr.append(avoid)
                if ridx==0:repeat=OxmlElement('w:tblHeader');pr.append(repeat)
                for j,(cell,value) in enumerate(zip(row.cells,values)):
                    cell.width=Inches(widths[j]);cell.vertical_alignment=1
                    pp=cell.paragraphs[0];pp.paragraph_format.space_after=Pt(4);pp.paragraph_format.space_before=Pt(4)
                    inline(pp,value)
                    for run in pp.runs:run.font.size=Pt(9)
                    tcpr=cell._tc.get_or_add_tcPr();borders=OxmlElement('w:tcBorders')
                    for side in ['top','left','bottom','right']:
                        edge=OxmlElement('w:'+side);edge.set(qn('w:val'),'single');edge.set(qn('w:sz'),'4');edge.set(qn('w:color'),'D9D9D9');borders.append(edge)
                    tcpr.append(borders)
                    margins=OxmlElement('w:tcMar')
                    for side in ['top','left','bottom','right']:
                        edge=OxmlElement('w:'+side);edge.set(qn('w:w'),'80');edge.set(qn('w:type'),'dxa');margins.append(edge)
                    tcpr.append(margins)
                    shade=OxmlElement('w:shd');shade.set(qn('w:fill'),'193B50' if ridx==0 else 'F2F5F7' if ridx%2==0 else 'FFFFFF');tcpr.append(shade)
                    if ridx==0:
                        for run in pp.runs:run.bold=True;run.font.color.rgb=RGBColor(255,255,255)
            doc.add_paragraph().paragraph_format.space_after=Pt(2)
            continue
        if line.startswith('# '):doc.add_paragraph(line[2:],'Title');i+=1;continue
        if line.startswith('## '):doc.add_paragraph(line[3:],'Heading 1');i+=1;continue
        if line.startswith('### '):doc.add_paragraph(line[4:],'Heading 2');i+=1;continue
        if re.match(r'^\d+\. ',line):
            p=doc.add_paragraph(style='List Number');inline(p,re.sub(r'^\d+\. ','',line));i+=1;continue
        if line.startswith('- '):
            p=doc.add_paragraph(style='List Bullet');inline(p,line[2:]);i+=1;continue
        text=line;i+=1
        while i<len(lines) and lines[i].strip() and not lines[i].startswith(('#','|','```','- ')):
            text+=' '+lines[i];i+=1
        inline(doc.add_paragraph(),text)
    doc.core_properties.title=doc.paragraphs[0].text
    doc.core_properties.author='Library study team'
    # Remove inherited decorative title rules from the base Word template.
    for tree in [doc._element, doc.styles.element]:
        for border in list(tree.iter(qn('w:pBdr'))):
            border.getparent().remove(border)
    output=source.with_suffix('.docx');doc.save(output);print(output)


if __name__=='__main__':
    for name in ['RESEARCH_REPORT.md','ENGINEERING_GUIDE.md']:build(ROOT/name)
